import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1A5276"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="0" w:color="BDC3C7"/><w:top w:val="single" w:sz="6" w:space="0" w:color="BDC3C7"/><w:right w:val="single" w:sz="6" w:space="0" w:color="BDC3C7"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="BDC3C7"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            
    for r_idx, row in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = "F2F5F8" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.0)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure_placeholder(doc, fig_num, caption, img_path=None):
    if img_path and os.path.exists(img_path):
        try:
            doc.add_picture(img_path, width=Inches(6.2))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass
    else:
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F4F6F7")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="dashed" w:sz="12" w:space="0" w:color="7F8C8D"/><w:top w:val="dashed" w:sz="12" w:space="0" w:color="7F8C8D"/><w:right w:val="dashed" w:sz="12" w:space="0" w:color="7F8C8D"/><w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="7F8C8D"/></w:tcBorders>')
        tcPr.append(borders)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[ Figure {fig_num} Screenshot: {caption} ]\n(Execution Output)")
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(3)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(f"Figure {fig_num} - {caption}")
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(9.0)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(0x56, 0x65, 0x73)

def create_base_doc(title, subtitle):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_t = p_title.add_run(title)
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(17)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    run_s = p_sub.add_run(subtitle)
    run_s.font.name = "Calibri"
    run_s.font.size = Pt(12)
    run_s.font.italic = True
    run_s.font.color.rgb = RGBColor(0x56, 0x65, 0x73)
    
    return doc

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
        run.font.size = Pt(12.5)
        run.bold = True
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    for run in h.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x28, 0x74, 0xA6)
        run.font.size = Pt(10.5)
        run.bold = True
    return h

def add_p(doc, text, bold_prefix="", space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10)
        r_b.bold = True
        r_b.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
    r_t = p.add_run(text)
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(10)
    r_t.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    return p

def generate_practical_15(output_path, img_map=None):
    if img_map is None:
        img_map = {}
        
    doc = create_base_doc(
        "MLOps Laboratory Manual - Practical 15",
        "Detect Data Drift and Trigger Model Retraining"
    )
    
    # Metadata Table
    headers = ["Student Name", "GitHub User", "Project", "Date"]
    rows = [["Bhanu Teja Velpuri", "BhanuTeja33", "Irrigation Need Prediction", "28-Oct-2026"]]
    add_styled_table(doc, headers, rows)
    
    # 1. Aim
    add_heading_1(doc, "1. Aim")
    add_p(doc, "To extend the production MLOps lifecycle of the Irrigation Need Prediction system by establishing a baseline feature reference, continuously monitoring incoming field telemetry for statistical data drift using two-sample Kolmogorov-Smirnov (KS) tests, and triggering candidate model retraining and quality gate evaluation only when significant data drift is detected and verified ground-truth labels are available.")
    
    # 2. Learning Objectives
    add_heading_1(doc, "2. Learning Objectives")
    objectives = [
        ("1. ML-Data Observability: ", "Understand why deployed ML serving APIs must be continuously monitored for shifts in input feature distributions beyond traditional system latency and error tracking."),
        ("2. Drift Taxonomy: ", "Distinguish data drift (covariate shift), concept drift (relationship shift), model performance drift (metric degradation), and normal statistical sampling variation."),
        ("3. Fixed Reference Baseline: ", "Create an immutable reference dataset representing the exact training distribution tied to the currently deployed model version."),
        ("4. Batch Simulation & Telemetry: ", "Simulate recent production batches representing both stable operating conditions and extreme environmental heatwave/drought drift shifts."),
        ("5. Statistical Hypothesis Testing: ", "Apply the non-parametric two-sample Kolmogorov-Smirnov (KS) test to evaluate distribution divergence across all 4 numerical telemetry features."),
        ("6. Machine-Readable Drift Auditing: ", "Generate structured JSON drift reports recording feature-level KS statistics, p-values, mean shifts, and overall drift classification."),
        ("7. Conditional Retraining Triggers: ", "Enforce a defensive trigger rule: prevent wasteful scheduled retraining when data are stable; initiate candidate retraining only when drift is statistically confirmed."),
        ("8. Ground-Truth Dependency: ", "Understand that unsupervised drift detection flags covariate shift, but supervised model retraining strictly requires labeled ground-truth outcomes."),
        ("9. Candidate vs. Production Isolation: ", "Train retrained models into isolated candidate artifacts (irrigation_model_candidate.pkl) without prematurely overwriting the active production model."),
        ("10. Quality Gate Enforcement: ", "Apply strict accuracy thresholds (>= 0.85) to candidate models before allowing promotion to production, safeguarding against regression."),
        ("11. End-to-End GitOps Traceability: ", "Connect data drift detection and model promotion to the automated GitHub Actions CI/CD and Kubernetes release pipeline developed in Practical 14.")
    ]
    for pfx, desc in objectives:
        add_p(doc, desc, bold_prefix=pfx, space_after=2)
        
    # 3. Continuity from Practical 14
    add_heading_1(doc, "3. Continuity from Practical 14")
    add_p(doc, "In Practical 14, the Irrigation Need Prediction microservice was integrated into an end-to-end GitHub Actions CI/CD pipeline, automatically publishing Docker images to GHCR and deploying them to Kubernetes via a self-hosted runner. Practical 15 enhances this cloud-native deployment with ML-level data monitoring. The historical dataset (field_data.csv), model serving API (app.py), Dockerfile, and CI/CD workflow remain intact; Practical 15 establishes the upstream intelligence deciding when model retraining should occur.")
    add_callout(doc, "Continuity Rule: Practical 15 does not automatically replace the production model merely because drift is detected. Drift serves as an automated trigger to retrain and evaluate; promotion still strictly requires passing the quality gate.", "CONTINUITY RULE")
    
    # 4. Drift Taxonomy
    add_heading_1(doc, "4. Drift Taxonomy: Data Drift vs. Concept Drift vs. Performance Drift")
    drift_headers = ["Drift Category", "Definition / Mechanism", "Irrigation Need Prediction Example", "Detection Method"]
    drift_rows = [
        ["Data Drift (Covariate Shift)", "Distribution of input features P(X) changes over time while relationship P(Y|X) remains steady.", "Severe summer drought causes average soil moisture to drop by 18% and temperature to spike by 8°C.", "Two-sample Kolmogorov-Smirnov (KS) test comparing reference vs. current features."],
        ["Concept Drift", "The mathematical mapping P(Y|X) between input features and target changes.", "A new drought-tolerant crop variety is planted, meaning crops survive lower moisture without needing irrigation.", "Requires ground-truth labels and residual/prediction error distribution tracking."],
        ["Performance Drift", "Observed model evaluation metrics (Accuracy, F1-Score) degrade over production time.", "Model accuracy drops from 88% to 68% on recently harvested crops.", "Requires delayed ground-truth feedback (post-harvest field outcome verification)."],
        ["Normal Sampling Variation", "Minor statistical fluctuation occurring naturally due to finite sample size.", "Day-to-day humidity variation of ±2% due to passing cloud cover.", "Statistical significance thresholds (p < 0.05) prevent false alert triggers."]
    ]
    add_styled_table(doc, drift_headers, drift_rows)
    
    # 5. Architecture
    add_heading_1(doc, "5. End-to-End Drift Detection & Retraining Lifecycle Architecture")
    add_p(doc, "The automated drift detection and candidate retraining workflow follows a defensive, multi-stage lifecycle:")
    
    arch_ascii = (
        "+-----------------------------------------------------------------------------------------+\n"
        "|                 Field IoT Telemetry Streams / Recent Production Batches                 |\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                                             |\n"
        "                                             v\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "|  STAGE 1: Statistical Drift Detection (Kolmogorov-Smirnov Two-Sample Test)               |\n"
        "|  - Baseline: reference_data.csv (Historical Training Distribution)                     |\n"
        "|  - Sample: current_data.csv (Recent Telemetry Batch, N=150)                             |\n"
        "|  - Test: ks_2samp(reference[feature], current[feature]) for all 4 features              |\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                                             |\n"
        "                                             v\n"
        "                             [Is overall_drift == True?]\n"
        "                             /                         \\\n"
        "                      [NO / Stable]              [YES / Drifted]\n"
        "                            |                           |\n"
        "                            v                           v\n"
        "+---------------------------------------+   +---------------------------------------------+\n"
        "| Action: CONTINUE_MONITORING           |   | Action: TRIGGER_RETRAINING                  |\n"
        "| - Log drift_report.json               |   | - Verify recent ground-truth labels exist   |\n"
        "| - Retraining SKIPPED                  |   |   (current_labeled_data.csv)                |\n"
        "| - Active production model preserved   |   +---------------------------------------------+\n"
        "+---------------------------------------+                           |\n"
        "                                                                    v\n"
        "                                            +---------------------------------------------+\n"
        "                                            |  STAGE 2: Candidate Model Retraining        |\n"
        "                                            |  - Merge historical + recent labeled data   |\n"
        "                                            |  - Train StandardScaler + LogisticRegression|\n"
        "                                            |  - Save: irrigation_model_candidate.pkl     |\n"
        "                                            |  - Record: retraining_metrics.json          |\n"
        "                                            +---------------------------------------------+\n"
        "                                                                    |\n"
        "                                                                    v\n"
        "                                                    [Quality Gate: Accuracy >= 0.85?]\n"
        "                                                     /                             \\\n"
        "                                              [FAIL / < 0.85]               [PASS / >= 0.85]\n"
        "                                                     |                              |\n"
        "                                                     v                              v\n"
        "                                        +-------------------------+   +-------------------+\n"
        "                                        | Candidate Rejected      |   | Candidate Promoted|\n"
        "                                        | - Retain old model      |   | - Backup old model|\n"
        "                                        | - Alert ML engineer     |   | - Overwrite prod  |\n"
        "                                        +-------------------------+   +-------------------+\n"
        "                                                                                |\n"
        "                                                                                v\n"
        "                                                                      +-------------------+\n"
        "                                                                      | Practical-14 CI/CD|\n"
        "                                                                      | - Docker build/tag|\n"
        "                                                                      | - Push to GHCR    |\n"
        "                                                                      | - Rolling rollout |\n"
        "                                                                      +-------------------+"
    )
    add_code_block(doc, arch_ascii)
    
    # 6. Step-by-Step Procedure
    add_heading_1(doc, "6. Step-by-Step Implementation Procedure & Verification")
    
    # Step 1
    add_heading_2(doc, "Step 1: Create the Reference Dataset Baseline")
    add_p(doc, "Extract the four numerical telemetry features from field_data.csv to establish the baseline distribution tied to the approved model:")
    add_code_block(doc, "python drift/prepare_reference.py\n# Output:\nReference dataset created successfully.\nRows: 400\nFile: .../drift/reference_data.csv")
    add_figure_placeholder(doc, 1, "Terminal execution of prepare_reference.py generating baseline reference_data.csv (400 rows)", img_map.get("fig1"))
    
    # Step 2
    add_heading_2(doc, "Step 2: Simulate Stable Telemetry & Verify No-Drift Case")
    add_p(doc, "Generate a recent production batch with standard environmental noise and run the drift detector:")
    add_code_block(doc, "python drift/simulate_current_data.py --mode stable\npython drift/detect_drift.py\n# Output:\nOverall Drift Detected: False\nFlagged Features (0): []\nAction: CONTINUE_MONITORING")
    add_figure_placeholder(doc, 2, "Drift detection in stable mode confirming 0 flagged features and CONTINUE_MONITORING", img_map.get("fig2"))
    
    # Step 3
    add_heading_2(doc, "Step 3: Verify Defensive Retraining Gate on Stable Data")
    add_p(doc, "Attempt to trigger candidate retraining when overall_drift is false. The script defensive guard immediately halts execution:")
    add_code_block(doc, "python drift/retrain_candidate.py\n# Output:\n============================================================\nDRIFT STATUS: NO OVERALL DRIFT DETECTED.\nRetraining skipped - preserving existing production model.\n============================================================")
    add_figure_placeholder(doc, 3, "Retraining guard preventing unnecessary compute execution when data are stable", img_map.get("fig3"))
    
    # Step 4
    add_heading_2(doc, "Step 4: Simulate Environmental Heatwave / Drought Drift")
    add_p(doc, "Simulate an extreme drought condition where soil moisture drops by 18%, temperature increases by 8°C, and humidity drops by 15%:")
    add_code_block(doc, "python drift/simulate_current_data.py --mode drift\n# Output:\nCurrent telemetry batch generated in 'drift' mode.\nRows: 150\nFile: .../drift/current_data.csv\nRecent labeled ground-truth dataset saved: .../drift/current_labeled_data.csv")
    add_figure_placeholder(doc, 4, "Simulation of extreme heatwave telemetry batch and ground-truth generation", img_map.get("fig4"))
    
    # Step 5
    add_heading_2(doc, "Step 5: Execute Kolmogorov-Smirnov Statistical Drift Detection")
    add_p(doc, "Run detect_drift.py against the drifted batch. All 4 features exhibit p-values < 0.05, triggering the overall drift alert:")
    add_code_block(doc, "python drift/detect_drift.py\n# Output:\n============================================================\nDATA DRIFT DETECTION REPORT\n============================================================\nOverall Drift Detected: True\nFlagged Features (4): ['soil_moisture', 'temperature', 'humidity', 'rainfall_forecast_mm']\nAction: TRIGGER_RETRAINING\nReport File: .../drift/drift_report.json")
    add_figure_placeholder(doc, 5, "Kolmogorov-Smirnov statistical test flagging 4 drifted features and triggering retraining", img_map.get("fig5"))
    
    # Step 6
    add_heading_2(doc, "Step 6: Retrain Candidate Model on Augmented Telemetry")
    add_p(doc, "Because drift is verified and recent ground truth is available, retrain_candidate.py merges historical and recent data, training irrigation_model_candidate.pkl:")
    add_code_block(doc, "python drift/retrain_candidate.py\n# Output:\n============================================================\nDRIFT DETECTED -> INITIATING CANDIDATE MODEL RETRAINING\n============================================================\nRetraining completed successfully.\n{\n  \"project\": \"Irrigation Need Prediction\",\n  \"candidate_model_path\": \"irrigation_model_candidate.pkl\",\n  \"candidate_accuracy\": 0.8545,\n  \"candidate_precision\": 0.84,\n  \"candidate_recall\": 0.84,\n  \"candidate_f1\": 0.84,\n  \"training_rows\": 440,\n  \"test_rows\": 110,\n  \"historical_rows\": 400,\n  \"recent_labeled_rows\": 150,\n  \"total_rows\": 550\n}\nCandidate model saved to: .../drift/irrigation_model_candidate.pkl")
    add_figure_placeholder(doc, 6, "Candidate model retrained successfully with 85.45% accuracy on augmented dataset", img_map.get("fig6"))
    
    # Step 7
    add_heading_2(doc, "Step 7: Evaluate Quality Gate and Promote Candidate")
    add_p(doc, "Evaluate the candidate model against the strict minimum accuracy threshold (>= 0.85). The candidate achieves 85.45%, passing the gate and promoting to production:")
    add_code_block(doc, "python drift/approve_candidate.py\n# Output:\n============================================================\nQUALITY GATE EVALUATION FOR CANDIDATE MODEL\n============================================================\nCandidate Model Accuracy: 0.8545 (85.45%)\nRequired Threshold      : 0.8500 (85.00%)\nExisting production model backed up to: irrigation_model_backup.pkl\n------------------------------------------------------------\nQUALITY GATE: PASS\nCandidate model successfully promoted to production: .../irrigation_model.pkl\nSanity Check Inference Test: soil=20, temp=38, hum=25, rain=0 -> Prediction: 1\n------------------------------------------------------------")
    add_figure_placeholder(doc, 7, "Quality gate evaluation confirming PASS (85.45% >= 85.00%) and promoting model", img_map.get("fig7"))
    
    # Step 8
    add_heading_2(doc, "Step 8: Execute Regression Test Suite on Promoted Model")
    add_p(doc, "Run the automated unittest suite to guarantee that API endpoints, feature schemas, and validation contracts remain 100% functional:")
    add_code_block(doc, "python -m unittest test_app.py test_irrigation_ml.py\n# Output:\n2026-10-28 15:14:40 | INFO | irrigation-api | event=model_loaded model_file=irrigation_model.pkl\nRan 11 tests in 0.074s - OK")
    add_figure_placeholder(doc, 8, "Unit test suite execution verifying zero regressions on newly promoted model (11/11 tests OK)", img_map.get("fig8"))
    
    # Step 9
    add_heading_2(doc, "Step 9: Version-Control Drift Scripts & Repository Synchronization")
    add_p(doc, "Commit the drift detection suite, baseline references, and retraining artifacts to GitHub to maintain full GitOps reproducibility:")
    add_code_block(doc, "git add drift/*.py generate_practical_15_docx.py\ngit commit -m \"Add Practical 15 data drift detection and candidate retraining pipeline\"\ngit push origin main")
    add_figure_placeholder(doc, 9, "GitHub repository tree showing drift suite and document generator committed to main branch", img_map.get("fig9"))

    # Step 10
    add_heading_2(doc, "Step 10: Trigger Automated Packaging & GHCR Publishing")
    add_p(doc, "The Git push triggers the GitHub Actions CI/CD workflow (build-test-publish), verifying the newly retrained model within the container test suite and pushing to GHCR in 1m 12s:")
    add_figure_placeholder(doc, 10, "GitHub Actions workflow run #12 executing build-test-publish with 100% success", img_map.get("fig10"))
    add_figure_placeholder(doc, 11, "GitHub Actions CI detailed step breakdown confirming model training, container health, and GHCR publishing", img_map.get("fig11"))
    
    # 7. Statistical Findings Table
    add_heading_1(doc, "7. Feature-Level Statistical Drift Analysis Table")
    stat_headers = ["Feature Name", "Reference Mean", "Current Mean", "Delta Mean", "KS Statistic", "p-value", "Drift Detected?"]
    stat_rows = [
        ["soil_moisture", "44.47 %", "26.75 %", "-17.73 %", "0.3567", "0.000000", "YES (Flagged)"],
        ["temperature", "27.41 °C", "35.77 °C", "+8.37 °C", "0.3300", "0.000000", "YES (Flagged)"],
        ["humidity", "57.03 %", "44.39 %", "-12.65 %", "0.2333", "0.000011", "YES (Flagged)"],
        ["rainfall_forecast_mm", "16.07 mm", "12.03 mm", "-4.04 mm", "0.2083", "0.000128", "YES (Flagged)"]
    ]
    add_styled_table(doc, stat_headers, stat_rows)
    
    # 8. Observation Table
    add_heading_1(doc, "8. Experimental Observation Table")
    obs_headers = ["Experiment Stage", "Tested Condition", "Expected Observation", "Observed Result"]
    obs_rows = [
        ["Baseline Establishment", "Historical field_data.csv", "400 rows extracted with 4 features", "reference_data.csv generated (400 rows)"],
        ["Stable Current Batch", "Small random noise (mode=stable)", "p-values > 0.05; overall_drift = false", "0 features flagged; action=CONTINUE_MONITORING"],
        ["Defensive Retraining Gate", "retrain_candidate on stable data", "Retraining blocked early", "Execution halted; production model preserved"],
        ["Drifted Batch", "Synthetic drought shift (mode=drift)", "Features shift significantly", "soil_moisture -17.7%, temp +8.4°C"],
        ["Statistical KS Test", "ks_2samp across all features", "p-values < 0.05 on >=2 features", "All 4 features flagged (p < 0.0002); overall_drift = true"],
        ["Candidate Retraining", "Augmented dataset (400 + 150)", "Candidate trained & evaluated separately", "Accuracy: 85.45%, F1: 0.84; candidate saved in drift/"],
        ["Quality Gate Evaluation", "Accuracy threshold >= 0.85", "Promote only if candidate meets threshold", "PASS: 85.45% >= 85.00%; promoted to irrigation_model.pkl"],
        ["Model Promotion Sanity", "joblib load & inference", "Model loads and predicts correctly", "Verified: [20, 38, 25, 0] -> IRRIGATE (code: 1)"],
        ["Regression Test Suite", "test_app.py & test_irrigation_ml.py", "All 11 automated test cases pass", "11/11 tests passed in 0.074s with HTTP 200/400/404 verified"]
    ]
    add_styled_table(doc, obs_headers, obs_rows)
    
    # 9. Key Commands Quick Reference
    add_heading_1(doc, "9. Key Commands Quick Reference Table")
    cmd_headers = ["Operational Step", "PowerShell Command"]
    cmd_rows = [
        ["Extract Reference Baseline", "python drift/prepare_reference.py"],
        ["Simulate Stable Telemetry", "python drift/simulate_current_data.py --mode stable"],
        ["Simulate Drifted Telemetry", "python drift/simulate_current_data.py --mode drift"],
        ["Execute Statistical Drift Test", "python drift/detect_drift.py"],
        ["Retrain Candidate Model", "python drift/retrain_candidate.py"],
        ["Evaluate Quality Gate & Promote", "python drift/approve_candidate.py"],
        ["Run Test Suite on Promoted Model", "python -m unittest test_app.py test_irrigation_ml.py"]
    ]
    add_styled_table(doc, cmd_headers, cmd_rows)
    
    # 10. Troubleshooting
    add_heading_1(doc, "10. Troubleshooting Common Data Drift & Retraining Issues")
    tb_headers = ["Observed Issue", "Root Cause", "Suggested Fix"]
    tb_rows = [
        ["FileNotFoundError: reference_data.csv", "Baseline preparation was skipped", "Execute python drift/prepare_reference.py first"],
        ["Missing required feature columns", "Dataset column headers differ from feature list", "Verify exact names: soil_moisture, temperature, humidity, rainfall_forecast_mm"],
        ["Retraining skipped unexpectedly", "overall_drift is false in drift_report.json", "Normal behavior on stable data; run simulate_current_data.py --mode drift to demonstrate shift"],
        ["current_labeled_data.csv missing", "Ground-truth labels not yet generated", "Run simulate_current_data.py which automatically saves labeled telemetry batch"],
        ["Quality Gate: FAIL", "Candidate accuracy < 0.85 threshold", "Inspect training distribution, balance classes, or adjust threshold before promotion"],
        ["Model loads but test suite fails", "Feature ordering or preprocessing mismatch", "Ensure StandardScaler is embedded inside the Pipeline alongside the model"]
    ]
    add_styled_table(doc, tb_headers, tb_rows)
    
    # 11. Viva Questions
    add_heading_1(doc, "11. Viva & Discussion Questions with Answers")
    viva = [
        ("Q1: What is data drift and why is it critical in MLOps?",
         "Data drift (covariate shift) occurs when the statistical distribution of input features changes over time compared to the data used during model training. In production, models can continue serving predictions without syntax errors while producing increasingly inaccurate inferences because the real-world population has fundamentally shifted."),
        ("Q2: How does data drift differ from concept drift and performance drift?",
         "Data drift involves changes strictly in input features P(X). Concept drift occurs when the fundamental relationship between inputs and outputs P(Y|X) changes. Performance drift is the observed degradation of evaluation metrics (e.g. accuracy drop), which requires labeled post-inference ground truth to calculate."),
        ("Q3: Why must a fixed reference dataset be established?",
         "Without a fixed reference tied to the currently deployed model version, the comparison baseline continuously moves. Comparing recent data against another recent batch would mask gradual, long-term environmental distribution shifts."),
        ("Q4: What statistical test is used in this practical and why?",
         "The two-sample Kolmogorov-Smirnov (KS) test (scipy.stats.ks_2samp) is used. It is a non-parametric test that evaluates whether two continuous empirical distributions differ significantly, without making normality assumptions about the features."),
        ("Q5: What does a p-value < 0.05 indicate in this drift detection system?",
         "A p-value below 0.05 indicates strong statistical evidence to reject the null hypothesis (that both samples originate from the identical distribution), confirming that the feature distribution has shifted significantly."),
        ("Q6: Why can unlabeled production inputs be used for drift detection but not retraining?",
         "Drift detection compares input feature vectors (soil moisture, temperature) which are readily available from incoming API telemetry. Supervised model retraining strictly requires ground-truth labels (whether the field actually required irrigation) to compute loss gradients and calculate validation accuracy."),
        ("Q7: Why is a retrained model saved as a candidate model first?",
         "According to defensive engineering principles, retrained models must never immediately replace production models. Saving as a candidate (irrigation_model_candidate.pkl) enables independent evaluation, offline validation, and quality-gate verification before promotion."),
        ("Q8: What is the role of the model quality gate in Practical 15?",
         "The quality gate acts as an automated safety filter. It verifies that the candidate model achieves at least 85% accuracy on holdout test data. If the candidate regresses due to noisy recent labels, it is rejected, preserving the existing stable production model."),
        ("Q9: Why are unittests rerun after model promotion?",
         "To guarantee that the newly serialized pipeline maintains full operational compatibility with the Flask serving application, feature dictionary parsers, and serialization libraries before triggering Docker containerization."),
        ("Q10: How does Practical 15 link to the Kubernetes CI/CD pipeline from Practical 14?",
         "Practical 15 acts as the intelligent trigger mechanism. When drift is detected, retrained, and approved, the updated model file is committed to Git. That Git commit automatically triggers the Practical 14 GitHub Actions CI/CD workflow, packaging the new model into a Docker image, publishing to GHCR, and rolling out to Kubernetes with zero downtime.")
    ]
    for q, a in viva:
        add_p(doc, a, bold_prefix=q + "\n", space_after=3)
        
    # 12. Result
    add_heading_1(doc, "12. Result")
    add_callout(doc, "Practical 15 was completed successfully. A statistical data drift detection and controlled retraining pipeline was established for the Irrigation Need Prediction system. Two-sample Kolmogorov-Smirnov tests successfully distinguished stable telemetry (0 flagged features) from extreme drought drift (4 flagged features, p < 0.0002). Retraining was executed conditionally on detected drift, and the resulting candidate model achieved 85.45% accuracy, passing the quality gate and promoting safely to production with 100% test suite verification.", "PRACTICAL RESULT: SUCCESS")
    
    # 13. Conclusion
    add_heading_1(doc, "13. Conclusion")
    add_p(doc, "This practical completes the machine learning feedback loop in the MLOps lifecycle. By moving beyond reactive system-level monitoring (CPU, latency) to proactive data-level statistical surveillance, the system detects distribution decay before silent prediction failures impact agricultural operations. The implementation of baseline references, hypothesis testing, candidate isolation, and automated quality gates ensures that model retraining is purposeful, mathematically justified, and enterprise-safe.")
    
    # Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Document saved successfully: {output_path}")

if __name__ == "__main__":
    extracted_dir = r"C:\Users\velpu\OneDrive\Desktop\mlops project\screenshots\p15_extracted"
    img_map = {
        "fig1": os.path.join(extracted_dir, "image_01.png"),
        "fig2": os.path.join(extracted_dir, "image_02.png"),
        "fig3": os.path.join(extracted_dir, "image_02.png"),
        "fig4": os.path.join(extracted_dir, "image_03.png"),
        "fig5": os.path.join(extracted_dir, "image_03.png"),
        "fig6": os.path.join(extracted_dir, "image_04.png"),
        "fig7": os.path.join(extracted_dir, "image_04.png"),
        "fig8": os.path.join(extracted_dir, "image_05.png"),
        "fig9": os.path.join(extracted_dir, "image_06.png"),
        "fig10": os.path.join(extracted_dir, "image_07.png"),
        "fig11": os.path.join(extracted_dir, "image_08.png")
    }
    out_dir = r"C:\Users\velpu\OneDrive\Desktop\mlops project\Practical_Word_Documents"
    out_file = os.path.join(out_dir, "PRACTICAL-15_Detect_Data_Drift_and_Retrain_Model.docx")
    generate_practical_15(out_file, img_map)
    
    root_file = r"C:\Users\velpu\OneDrive\Desktop\mlops project\PRACTICAL-15_Detect_Data_Drift_and_Retrain_Model_Final.docx"
    generate_practical_15(root_file, img_map)
